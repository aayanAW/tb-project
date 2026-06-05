#!/usr/bin/env python3
"""Convert METHODOLOGY_DOCUMENT.md to a professionally formatted .docx file."""

import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

doc = Document()

# -- Page setup --
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# -- Style setup --
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for level, (size, color) in enumerate([
    (24, '1B4F72'), (18, '1B4F72'), (14, '2E75B6'), (12, '2E75B6')
], start=1):
    hstyle = doc.styles[f'Heading {level}']
    hstyle.font.name = 'Calibri'
    hstyle.font.size = Pt(size)
    hstyle.font.color.rgb = RGBColor(*bytes.fromhex(color))
    hstyle.font.bold = True
    hstyle.paragraph_format.space_before = Pt(18 if level <= 2 else 12)
    hstyle.paragraph_format.space_after = Pt(6)

# -- Read the markdown --
with open('/Users/aayanalwani/tb project/mce3r_stochastic/results/METHODOLOGY_DOCUMENT.md') as f:
    md = f.read()

# -- Title page --
for _ in range(6):
    doc.add_paragraph('')

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Mce3R Operator Discovery &\nStochastic Gene Expression Modeling')
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)
run.font.bold = True
run.font.name = 'Calibri'

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Complete Methodology Document for Regeneron STS')
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
run.font.name = 'Calibri'

doc.add_paragraph('')
author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = author.add_run('Aayan Alwani\nBalazsi Lab, Stony Brook University')
run.font.size = Pt(13)
run.font.name = 'Calibri'

doc.add_page_break()

# -- Parse markdown and convert --
lines = md.split('\n')
i = 0
in_code_block = False
code_lines = []
in_table = False
table_rows = []

def add_code_block(doc, code_text):
    """Add a formatted code block."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    # Add shading
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8" w:val="clear"/>')
    p.paragraph_format.element.get_or_add_pPr().append(shading)

def add_table(doc, rows):
    """Add a formatted table."""
    if not rows or len(rows) < 2:
        return
    # Parse header and data
    def parse_row(row):
        cells = [c.strip() for c in row.strip('|').split('|')]
        return cells

    headers = parse_row(rows[0])
    # Skip separator row (row[1])
    data_rows = [parse_row(r) for r in rows[2:] if r.strip()]

    ncols = len(headers)
    table = doc.add_table(rows=1 + len(data_rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Grid Accent 1'

    # Header row
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(format_inline(h))
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.name = 'Calibri'
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Dark header background
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B4F72" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # Data rows
    for i_row, dr in enumerate(data_rows):
        for j, val in enumerate(dr):
            if j >= ncols:
                break
            cell = table.rows[i_row + 1].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(format_inline(val))
            run.font.size = Pt(10)
            run.font.name = 'Calibri'
            # Alternating row colors
            if i_row % 2 == 0:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EBF5FB" w:val="clear"/>')
                cell._tc.get_or_add_tcPr().append(shading)

    doc.add_paragraph('')  # spacing after table

def format_inline(text):
    """Strip markdown bold/italic markers for plain text."""
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text

def add_rich_paragraph(doc, text, style_name='Normal', is_bullet=False, bullet_level=0):
    """Add a paragraph with bold/italic/code formatting."""
    p = doc.add_paragraph(style=style_name)
    if is_bullet:
        p.style = doc.styles['List Bullet']
        if bullet_level > 0:
            p.style = doc.styles['List Bullet 2'] if bullet_level == 1 else doc.styles['List Bullet']

    # Parse inline formatting
    # Split on **bold**, *italic*, `code`
    pattern = r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*|`.+?`)'
    parts = re.split(pattern, text)

    for part in parts:
        if part.startswith('***') and part.endswith('***'):
            run = p.add_run(part[3:-3])
            run.font.bold = True
            run.font.italic = True
        elif part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = p.add_run(part[1:-1])
            run.font.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
        else:
            run = p.add_run(part)

    return p

# Skip the first two lines (title and subtitle, already on title page)
# Find where content starts (after first ---)
start = 0
for idx, line in enumerate(lines):
    if line.strip() == '---':
        start = idx + 1
        break

# If no --- found, skip first 4 lines
if start == 0:
    start = 4

i = start
while i < len(lines):
    line = lines[i]

    # Code blocks
    if line.strip().startswith('```'):
        if not in_code_block:
            in_code_block = True
            code_lines = []
            i += 1
            continue
        else:
            in_code_block = False
            add_code_block(doc, '\n'.join(code_lines))
            i += 1
            continue

    if in_code_block:
        code_lines.append(line)
        i += 1
        continue

    # Table rows
    if '|' in line and line.strip().startswith('|'):
        if not in_table:
            in_table = True
            table_rows = []
        table_rows.append(line)
        i += 1
        continue
    elif in_table:
        in_table = False
        add_table(doc, table_rows)
        table_rows = []
        # Don't increment, process current line

    # Horizontal rule
    if line.strip() == '---':
        # Add a thin line
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        pPr = p.paragraph_format.element.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="2E75B6"/></w:pBdr>')
        pPr.append(pBdr)
        i += 1
        continue

    # Headings
    if line.startswith('# ') and not line.startswith('## '):
        text = line[2:].strip()
        text = format_inline(text)
        h = doc.add_heading(text, level=1)
        i += 1
        continue
    if line.startswith('## '):
        text = line[3:].strip()
        text = format_inline(text)
        doc.add_heading(text, level=2)
        i += 1
        continue
    if line.startswith('### '):
        text = line[4:].strip()
        text = format_inline(text)
        doc.add_heading(text, level=3)
        i += 1
        continue

    # Bullet points
    if line.strip().startswith('- '):
        indent = len(line) - len(line.lstrip())
        level = 0 if indent < 2 else 1
        text = line.strip()[2:]
        add_rich_paragraph(doc, text, is_bullet=True, bullet_level=level)
        i += 1
        continue

    # Numbered lists
    m = re.match(r'^(\d+)\.\s+(.*)', line.strip())
    if m:
        text = m.group(2)
        add_rich_paragraph(doc, text, is_bullet=True)
        i += 1
        continue

    # Empty lines
    if not line.strip():
        i += 1
        continue

    # Normal paragraphs
    # Collect continuation lines
    para_text = line.strip()
    while i + 1 < len(lines):
        next_line = lines[i + 1]
        if (not next_line.strip() or
            next_line.startswith('#') or
            next_line.strip().startswith('- ') or
            next_line.strip().startswith('```') or
            next_line.strip() == '---' or
            re.match(r'^\d+\.', next_line.strip()) or
            (next_line.strip().startswith('|') and '|' in next_line)):
            break
        para_text += ' ' + next_line.strip()
        i += 1

    add_rich_paragraph(doc, para_text)
    i += 1

# Handle any remaining table
if in_table and table_rows:
    add_table(doc, table_rows)

# -- Save --
output_path = '/Users/aayanalwani/tb project/mce3r_stochastic/results/METHODOLOGY_DOCUMENT.docx'
doc.save(output_path)
print(f"Saved to: {output_path}")
