#!/usr/bin/env python3
"""Generate ENIGMA Short Proposal (5-page condensed version) as DOCX."""

from docx import Document
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

doc = Document()

# ── Page setup: 1-inch margins, US Letter ──
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

# ── Default style: Calibri 11pt, single-spaced, 6pt after ──
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
pf = style.paragraph_format
pf.space_before = Pt(0)
pf.space_after = Pt(6)
pf.line_spacing = 1.0

# Helper: set font on a run
def set_run_font(run, size=11, bold=False, italic=False, name='Calibri'):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

# Helper: add a paragraph with specific formatting
def add_para(text, size=11, bold=False, italic=False, alignment=None, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.0
    if alignment:
        p.alignment = alignment
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p

# Helper: add section header (bold 12pt)
def add_section_header(text, space_before=10):
    return add_para(text, size=12, bold=True, space_after=4, space_before=space_before)

# Helper: add subsection header (bold 11pt)
def add_subsection_header(text, space_before=8):
    return add_para(text, size=11, bold=True, space_after=3, space_before=space_before)

# Helper: add body text
def add_body(text, space_after=6):
    return add_para(text, size=11, space_after=space_after)

# Helper: add mixed-format paragraph
def add_mixed_para(parts, space_after=6, space_before=0, alignment=None):
    """parts is list of (text, size, bold, italic) tuples"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.0
    if alignment:
        p.alignment = alignment
    for text, size, bold, italic in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    return p

# ═══════════════════════════════════════════════════════════════
# TITLE
# ═══════════════════════════════════════════════════════════════
add_para(
    "The Persistence Regulatory Code: Genome-Wide Operator Asymmetry, Temporal Noise Memory, and Therapeutic Noise Quenching in Mycobacterium tuberculosis",
    size=14, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4
)

# SUBTITLE
add_para(
    "ENIGMA: Expression Noise In Gene-regulatory Mechanisms and Architecture",
    size=10, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=10
)

# ═══════════════════════════════════════════════════════════════
# ABSTRACT
# ═══════════════════════════════════════════════════════════════
add_section_header("Abstract", space_before=2)

add_body(
    "Antibiotic-tolerant persister cells drive tuberculosis treatment failure. Gene expression noise generates these cells, and cis-regulatory architecture shapes noise, but no study has systematically characterized how operator-level DNA architecture contributes to persistence-relevant noise across the M. tuberculosis genome. We present an integrated computational framework anchored by Mce3R, a repressor with an unprecedented 20.4-fold binding site asymmetry (Kd = 2.4 vs. 49 nM). We discover 1,442 candidate binding sites genome-wide, build a thermodynamically calibrated stochastic model (\u03c9 = 1.08), and show the asymmetric operator generates 19% more noise (CV = 0.187 vs. 0.157, p < 0.001)\u2014robust across autoregulation, four environments, and the full MCMC posterior. We extend to genome-wide asymmetry scoring of persistence genes, temporal noise dynamics and multi-generational lineage memory, and noise-quenching dose-response predictions. Results suggest operator asymmetry may represent a genome-encoded persistence strategy and identify quantitative therapeutic targets."
)

# ═══════════════════════════════════════════════════════════════
# 1. INTRODUCTION
# ═══════════════════════════════════════════════════════════════
add_section_header("1. Introduction")

add_body(
    "Tuberculosis kills 1.2 million people annually (WHO 2024). Treatment requires 6\u20139 months of daily antibiotics because phenotypically tolerant persister cells survive drug exposure and drive relapse (Balaban et al. 2019). Persisters arise from stochastic gene expression noise\u2014genetically identical cells randomly enter a drug-tolerant dormant state. Reducing the persister subpopulation could dramatically shorten treatment."
)

add_body(
    "Cis-regulatory architecture shapes expression noise (Chowdhury et al. 2021), multiple binding sites modulate noise depending on configuration (Lengyel & Morelli 2017), and noise in metabolic pathways directly produces M. tuberculosis persisters (Quigley & Lewis 2022). However, no study has systematically identified which specific architectural features of regulatory DNA generate persistence-relevant noise, whether these features are enriched at persistence-critical genes, or how their temporal properties affect persistence duration."
)

add_body(
    "The Mce3R repressor controls cholesterol metabolism genes essential for host survival (Santangelo et al. 2002). Its deletion increases persister frequency (Pandey et al. 2023). Cryo-EM revealed an asymmetric 123 bp operator with a 20.4-fold Kd difference between two binding sites (2.4 nM vs. 49 nM; Panagoda et al. 2024, PDB 9B7Y)\u2014the largest documented asymmetry in any TetR-family system. This system is uniquely suited for computational noise analysis because it has real, structurally characterized binding parameters."
)

add_body(
    "Three questions remain open: (1) Does this specific asymmetry produce quantifiably more noise than symmetric alternatives? (2) Is operator asymmetry enriched at persistence-critical genes genome-wide? (3) What are the temporal properties of asymmetry-driven noise, and how much symmetrization would reduce persistence? We address these through four aims: genome-wide binding site discovery and asymmetry scoring (Aim 1), thermodynamically calibrated stochastic modeling (Aim 2), comprehensive noise characterization across amplitude, temporal dynamics, lineage memory, and quenching (Aim 3), and evolutionary simulation of asymmetry under persistence selection (Aim 4, stretch). All results are computational predictions with explicit falsification criteria."
)

# ═══════════════════════════════════════════════════════════════
# 2. AIMS AND METHODS
# ═══════════════════════════════════════════════════════════════
add_section_header("2. Aims and Methods")

# Aim 1
add_mixed_para([
    ("Aim 1: Discover Binding Sites and Test Genome-Wide Persistence Enrichment", 11, True, False),
], space_after=3, space_before=4)

add_body(
    "Applied MEME de novo motif discovery (ZOOPS mode, width 6\u2013110 bp) to 200 bp upstream of yrbE3A orthologs from M. tuberculosis H37Rv, M. bovis, and M. marinum. Scanned the full H37Rv genome with FIMO (p < 1\u00d710\u207b\u2074), identifying 1,442 candidate Mce3R binding sites. For genome-wide analysis, we compute an asymmetry score for each site: |log\u2082(PWM_score_strong / PWM_score_weak)| for paired binding sites within 150 bp windows. Additional features include palindromicity index, spacer length, and information content. We cross-reference all sites against published Tn-seq persistence gene lists (DeJesus et al. 2017) and test enrichment using Mann-Whitney U and Fisher\u2019s exact tests. We then run calibrated Gillespie simulations on the top 20\u201350 highest-asymmetry operators to predict their noise output."
)

add_mixed_para([
    ("Expected result: ", 11, True, True),
    ("Persistence genes have significantly higher median asymmetry scores, suggesting M. tb deploys high-asymmetry operators at genes that benefit from noise-driven phenotypic diversification.", 11, False, False),
], space_after=6)

# Aim 2
add_mixed_para([
    ("Aim 2: Build Thermodynamically Calibrated Stochastic Model", 11, True, False),
], space_after=3, space_before=4)

add_body(
    "Modeled the operator as four states (unbound, strong-only, weak-only, both-bound) with 12 reactions in a Gillespie SSA (Numba JIT-accelerated). Kinetic parameters from published sources: k_on = 0.0167 nM\u207b\u00b9min\u207b\u00b9 (Stormo & Zhao 2010), mRNA t\u00bd = 9.5 min (Rustad et al. 2013), protein t\u00bd = 1500 min, k_translation = 0.5/min (Taniguchi et al. 2010). Simulated 50,000 cells per condition across four architectures: asymmetric native (Kd = 2.4/49 nM), symmetric control (Kd = 10.84 nM geometric mean), single-site, and unregulated. Calibrated binding energies via the Berg-von Hippel model (\u0394G_strong = \u221212.23, \u0394G_weak = \u221210.37 kcal/mol). Inferred cooperativity via MCMC (32 walkers \u00d7 5,000 steps): \u03c9 = 1.08 [95% CI: 0.20\u20135.89], indicating no strong cooperativity\u2014the sites bind independently, creating more distinct operator states and more noise."
)

# Aim 3
add_mixed_para([
    ("Aim 3: Characterize Noise Amplitude, Temporal Memory, and Quenching", 11, True, False),
], space_after=3, space_before=4)

add_mixed_para([
    ("Sub-aim 3a \u2014 Amplitude: ", 11, True, True),
    ("Measured CV across 3 architectures \u00d7 4 environments (baseline, cholesterol, acidic pH, host-like) \u00d7 2 models (single-species, two-species autoregulatory). Propagated MCMC posterior uncertainty (50 samples). Computed mutual information and calibrated persister fractions.", 11, False, False),
], space_after=6)

add_mixed_para([
    ("Sub-aim 3b \u2014 Temporal dynamics and lineage memory (NEW): ", 11, True, True),
    ("Record Gillespie time traces at 10-min intervals (1,500 points per cell, 1,000 cells per condition). Compute autocorrelation C(\u03c4) via FFT, fit exponential decay to extract \u03c4_c (noise memory duration). Compute dwell time distributions in persister state and power spectral density. For lineage tracking: simulate cell division at 1,500-min intervals (Mtb doubling time), daughter inherits protein/2 (binomial partitioning) + operator state. Track 100 lineage trees of 128 cells (7 generations) per architecture. Measure: generations until all descendants exit persister state.", 11, False, False),
], space_after=6)

add_mixed_para([
    ("Sub-aim 3c \u2014 Noise quenching dose-response (NEW): ", 11, True, True),
    ("Sweep Kd_weak from 49 \u2192 2.4 nM (15 steps), measuring CV and persister fraction at each. Sweep [Mce3R] from 10\u201310,000 nM. Fit sigmoidal dose-response, compute IC50_noise (symmetrization halfway point) and IC50_persistence (50% persister reduction). This produces the first quantitative dose-response prediction for noise quenching in a TB system.", 11, False, False),
], space_after=6)

# Aim 4
add_mixed_para([
    ("Aim 4 (Stretch): Evolutionary Simulation", 11, True, False),
], space_after=3, space_before=4)

add_body(
    "Fit polynomial surrogate CV(Kd_ratio, spacer, block) from sweep data. Wrap in genetic algorithm: population 200, 1,000 generations, mutable operator parameters. Two selection regimes: (A) persistence selection\u2014fitness = survival fraction under antibiotic pulse; (B) growth selection\u2014fitness = mean expression. Test whether asymmetry increases under persistence selection but not growth selection. Novel prediction: operator asymmetry is positively selected as a noise generator specifically under persistence pressure."
)

# ═══════════════════════════════════════════════════════════════
# 3. RESULTS AND SIGNIFICANCE
# ═══════════════════════════════════════════════════════════════
add_section_header("3. Results and Significance")

add_mixed_para([
    ("Aim 1: ", 11, True, False),
    ("FIMO top hit matched the known operator (p = 1.1 \u00d7 10\u207b\u00b9\u00b9), validating the computational pipeline. 1,442 candidate sites identified genome-wide, including novel targets near tgs1 (dormancy), espC (virulence), and mbtG (iron acquisition). Genome-wide asymmetry enrichment analysis pending\u2014if confirmed, this establishes a \u201cpersistence regulatory code\u201d encoded in operator DNA architecture.", 11, False, False),
], space_after=6)

add_mixed_para([
    ("Aim 2: ", 11, True, False),
    ("Berg-von Hippel calibration reproduced both Kd values exactly. MCMC inferred \u03c9 = 1.08 (no cooperativity)\u2014the two sites bind independently, creating four distinct operator states with different transcription rates. All 20 classified FIMO operators are graded repressors (Hill n < 2), not digital switches. The asymmetric architecture creates a broader repression transition region, allowing more intermediate expression states.", 11, False, False),
], space_after=6)

add_mixed_para([
    ("Aim 3a: ", 11, True, False),
    ("CV(asymmetric) = 0.187 vs. CV(symmetric) = 0.157 (p < 0.001, KS = 0.753, Cohen\u2019s d = \u22122.29). Bootstrap 95% CI for \u0394CV entirely above zero. Result holds in 100% of MCMC posterior samples, in the two-species autoregulatory model (CV = 0.200 vs. 0.197), and across all four environments. Calibrated persister fractions: 0.28% asymmetric vs. 0.09% symmetric (3.1\u00d7 enrichment baseline; 29.4\u00d7 under acidic pH). MI(asymmetric) = 0.28 bits < MI(symmetric) = 0.44 bits\u2014the asymmetric operator trades information for noise, consistent with bet-hedging (Veening et al. 2008). Context: CV = 0.187 is moderate (25% above typical regulated genes, 42% of hipA TA module). Operator asymmetry provides a constitutive noise floor, not the dominant noise source.", 11, False, False),
], space_after=6)

add_mixed_para([
    ("Aim 3b (predicted): ", 11, True, False),
    ("\u03c4_c(asymmetric) > \u03c4_c(symmetric)\u2014asymmetric operators generate noise with longer temporal memory because the strong site stays bound (slow k_off = 0.04/min) while the weak site flickers (fast k_off = 0.82/min), creating extended partial-repression episodes. Lineage tracking predicts asymmetric operator produces longer persistence lineages\u2014more generations until all descendants exit the persister state. Clinical implication: asymmetry may affect treatment duration, not just persister frequency.", 11, False, False),
], space_after=6)

add_mixed_para([
    ("Aim 3c (predicted): ", 11, True, False),
    ("Sigmoidal dose-response for CV vs. Kd_weak. IC50_noise provides a quantitative therapeutic target. Flentie et al. (2019) showed Mce3R-targeting compounds enhance antibiotics 16\u201350\u00d7; this model predicts the mechanism and required dose.", 11, False, False),
], space_after=6)

add_mixed_para([
    ("Therapeutic significance: ", 11, True, False),
    ("If noise drives persistence, reducing noise may reduce persisters. The project identifies a specific, quantifiable architectural feature (the Kd ratio) as a potential \u201canti-noise\u201d drug target\u2014not killing bacteria directly, but making them uniformly susceptible to existing antibiotics.", 11, False, False),
], space_after=6)

# ═══════════════════════════════════════════════════════════════
# 4. LIMITATIONS AND FALSIFICATION
# ═══════════════════════════════════════════════════════════════
add_section_header("4. Limitations and Falsification")

add_mixed_para([
    ("Limitations: ", 11, True, False),
    ("(1) All results are computational\u2014no experimental validation. (2) Single/two-gene model of a 14-gene operon. (3) Wide MCMC posterior for \u03c9. (4) Environments modeled as scalar multipliers. (5) CV difference is moderate; other noise sources likely larger. (6) \u03c4_c may be dominated by protein half-life. (7) Genome-wide enrichment depends on Tn-seq annotation quality.", 11, False, False),
], space_after=6)

add_mixed_para([
    ("Falsification criteria: ", 11, True, False),
    ("The hypothesis is falsified if (1) experimentally symmetrizing the operator shows no noise change, (2) upstream noise sources eliminate the CV difference, (3) strong cooperativity suppresses independent binding noise, (4) \u03c4_c shows no architecture dependence, or (5) persistence genes show no asymmetry enrichment.", 11, False, False),
], space_after=6)

# ═══════════════════════════════════════════════════════════════
# 5. REFERENCES
# ═══════════════════════════════════════════════════════════════
add_section_header("5. References")

references = [
    "1. WHO (2024). Global TB Report.",
    "2. Balaban et al. (2019). Nat Rev Microbiol 17:441.",
    "3. Elowitz et al. (2002). Science 297:1183.",
    "4. Chowdhury et al. (2021). Front Genet 12:698910.",
    "5. Lengyel & Morelli (2017). Phys Rev E 95:042412.",
    "6. Balazsi et al. (2011). Cell 144:910.",
    "7. Quigley & Lewis (2022). Microbiol Spectrum 10:e0094822.",
    "8. Santangelo et al. (2002). Microbiology 148:2997.",
    "9. Pandey et al. (2023). Res Microbiol 174:104082.",
    "10. Panagoda et al. (2024). ACS Chem Biol 19:2580.",
    "11. DeJesus et al. (2017). mBio 8:e02133.",
    "12. Gillespie (1977). J Phys Chem 81:2340.",
    "13. Stormo & Zhao (2010). Nat Rev Genet 11:751.",
    "14. Taniguchi et al. (2010). Science 329:533.",
    "15. Rustad et al. (2013). Nucleic Acids Res 41:509.",
    "16. Flentie et al. (2019). ACS Infect Dis 5:1239.",
    "17. Veening et al. (2008). Annu Rev Microbiol 62:193.",
]

for ref in references:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    # Hanging indent: left=0.3in, first_line=-0.3in
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    run = p.add_run(ref)
    set_run_font(run, size=10)

# ── Save ──
out_path = "/Users/aayanalwani/tb project/mce3r_stochastic/ENIGMA_Short_Proposal.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
print(f"Size: {os.path.getsize(out_path)} bytes ({os.path.getsize(out_path)/1024:.1f} KB)")

# Estimate word count
total_words = 0
for p in doc.paragraphs:
    total_words += len(p.text.split())
print(f"Word count: ~{total_words}")
# Rough page estimate: single-spaced Calibri 11pt ~ 500-550 words/page
print(f"Estimated pages: ~{total_words/520:.1f}")
