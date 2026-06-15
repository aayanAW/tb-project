#!/usr/bin/env python3
"""Generate Phase 4 and Phase 7 presentations using python-pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import os

# ── Colors ────────────────────────────────────────────────────────────────
TEAL = RGBColor(0x0D, 0x94, 0x88)
DARK_TEAL = RGBColor(0x0A, 0x6E, 0x64)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1E, 0x29, 0x3B)
GRAY = RGBColor(0x64, 0x74, 0x8B)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
LIGHT_GRAY = RGBColor(0xE2, 0xE8, 0xF0)
ACCENT_ORANGE = RGBColor(0xF9, 0x73, 0x16)
ACCENT_PURPLE = RGBColor(0x8B, 0x5C, 0xF6)

BASE = "/Users/aayanalwani/tb project/mce3r_stochastic"
FIGURES_DIR = os.path.join(BASE, "results", "figures")
EXT_FIGURES_DIR = os.path.join(BASE, "results", "extended_figures")


# ── Helpers ───────────────────────────────────────────────────────────────
def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(slide, x, y, w, h, text, size=16, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_bullets(slide, x, y, w, h, items, size=14, color=DARK, spacing=Pt(6)):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            run_b = p.add_run()
            run_b.text = item[0]
            run_b.font.bold = True
            run_b.font.size = Pt(size)
            run_b.font.color.rgb = color
            run_b.font.name = 'Calibri'
            run_n = p.add_run()
            run_n.text = item[1]
            run_n.font.size = Pt(size)
            run_n.font.color.rgb = color
            run_n.font.name = 'Calibri'
        else:
            p.text = item
            p.font.size = Pt(size)
            p.font.color.rgb = color
            p.font.name = 'Calibri'
        p.space_after = spacing
        pPr = p._p.get_or_add_pPr()
        buChar = etree.SubElement(pPr, qn('a:buChar'))
        buChar.set('char', '\u2022')
    return txBox


def add_image(slide, path, x, y, w=None, h=None):
    if not os.path.exists(path):
        # Placeholder box
        pw = w or 5
        ph = h or 3
        shape = add_rect(slide, x, y, pw, ph, LIGHT_GRAY)
        shape.text_frame.paragraphs[0].text = f"[{os.path.basename(path)}]"
        shape.text_frame.paragraphs[0].font.color.rgb = GRAY
        shape.text_frame.paragraphs[0].font.size = Pt(12)
        shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        return
    kwargs = {"image_file": path, "left": Inches(x), "top": Inches(y)}
    if w:
        kwargs["width"] = Inches(w)
    if h:
        kwargs["height"] = Inches(h)
    slide.shapes.add_picture(**kwargs)


def make_title_slide(prs, title, subtitle):
    slide = add_blank(prs)
    set_bg(slide, DARK)
    # Left teal accent bar
    add_rect(slide, 0, 0, 0.15, 7.5, TEAL)
    # Title
    add_text(slide, 1.0, 2.0, 11.0, 1.2, title, size=40, color=WHITE, bold=True,
             align=PP_ALIGN.LEFT, font_name='Trebuchet MS')
    # Accent line
    add_rect(slide, 1.0, 3.4, 4.0, 0.06, TEAL)
    # Subtitle
    add_text(slide, 1.0, 3.7, 11.0, 0.8, subtitle, size=22, color=LIGHT_GRAY,
             align=PP_ALIGN.LEFT)
    # Bottom bar
    add_rect(slide, 0, 7.1, 13.333, 0.4, TEAL)
    add_text(slide, 0.5, 7.12, 12.0, 0.35, "ENIGMA Project  |  Mce3R Stochastic Gene Regulation",
             size=12, color=WHITE, align=PP_ALIGN.LEFT)
    return slide


def make_header_bar(slide, title):
    """Add a teal header bar at top of a content slide."""
    add_rect(slide, 0, 0, 13.333, 0.9, TEAL)
    add_text(slide, 0.5, 0.1, 12.0, 0.7, title, size=28, color=WHITE, bold=True,
             font_name='Trebuchet MS')


def make_content_slide(prs, title):
    slide = add_blank(prs)
    set_bg(slide, LIGHT_BG)
    make_header_bar(slide, title)
    # Bottom bar
    add_rect(slide, 0, 7.1, 13.333, 0.4, DARK)
    add_text(slide, 0.5, 7.12, 6.0, 0.35, "ENIGMA Project", size=10, color=LIGHT_GRAY)
    return slide


def make_figure_slide(prs, title, img_path, caption):
    """Create a figure slide with the image filling most of the area."""
    slide = make_content_slide(prs, title)
    # Image: centered, large
    add_image(slide, img_path, x=1.0, y=1.1, w=11.3, h=5.2)
    # Caption below image
    add_text(slide, 0.5, 6.4, 12.3, 0.6, caption, size=13, color=GRAY,
             align=PP_ALIGN.CENTER)
    return slide


def make_summary_slide(prs, title, body_text):
    slide = add_blank(prs)
    set_bg(slide, DARK)
    add_rect(slide, 0, 0, 0.15, 7.5, TEAL)
    add_text(slide, 1.0, 1.5, 11.0, 1.0, title, size=36, color=WHITE, bold=True,
             font_name='Trebuchet MS')
    add_rect(slide, 1.0, 2.7, 3.5, 0.06, TEAL)
    add_text(slide, 1.0, 3.2, 11.0, 3.5, body_text, size=20, color=LIGHT_GRAY)
    add_rect(slide, 0, 7.1, 13.333, 0.4, TEAL)
    add_text(slide, 0.5, 7.12, 12.0, 0.35, "ENIGMA Project  |  Mce3R Stochastic Gene Regulation",
             size=12, color=WHITE, align=PP_ALIGN.LEFT)
    return slide


# ══════════════════════════════════════════════════════════════════════════
#  PHASE 4 PRESENTATION
# ══════════════════════════════════════════════════════════════════════════
def create_phase4():
    prs = new_prs()

    # ── Slide 1: Title ────────────────────────────────────────────────────
    make_title_slide(prs, "Phase 4: Manuscript Figures",
                     "ENIGMA Project \u2014 Visualizing the Results")

    # ── Slide 2: Figure Overview ──────────────────────────────────────────
    slide = make_content_slide(prs, "Figure Overview")
    items = [
        ("Fig 1 \u2014 Binding Site Architecture: ", "Circular genome plot with FIMO-scored operator sites"),
        ("Fig 2 \u2014 Protein Distributions: ", "Histograms + GMM fits for 4 regulatory conditions"),
        ("Fig 3 \u2014 Asymmetry Sweep: ", "CV and intermediate fraction vs Kd ratio"),
        ("Fig 4 \u2014 Sensitivity Analysis: ", "Tornado plot + heatmap of parameter perturbation effects"),
        ("Fig 5 \u2014 Experimental Validation: ", "Model predictions vs Santangelo 2009 data"),
        ("Fig 6 \u2014 Single-Cell Traces: ", "Individual cell time series for 3 conditions"),
        ("Fig 7 \u2014 BIC Model Comparison: ", "Grouped bar chart of GMM component selection"),
        ("Fig 8 \u2014 Methods Architecture: ", "Computational pipeline diagram across all phases"),
    ]
    add_bullets(slide, 0.6, 1.2, 12.0, 5.5, items, size=16, spacing=Pt(10))

    # ── Slide 3: Figure 1 ────────────────────────────────────────────────
    make_figure_slide(prs, "Figure 1: Binding Site Architecture",
                      os.path.join(FIGURES_DIR, "fig1_binding_sites.png"),
                      "Circular genome plot of H37Rv with predicted Mce3R binding sites. "
                      "Inset: top 15 sites ranked by FIMO score.")

    # ── Slide 4: Figure 2 ────────────────────────────────────────────────
    make_figure_slide(prs, "Figure 2: Protein Distributions",
                      os.path.join(FIGURES_DIR, "fig2_distributions.png"),
                      "2\u00d73 panel: histograms with GMM fits for 4 conditions (A\u2013D), "
                      "box plot of CV/Fano/bimodality, and density overlay.")

    # ── Slide 5: Figure 3 ────────────────────────────────────────────────
    make_figure_slide(prs, "Figure 3: Asymmetry Sweep",
                      os.path.join(FIGURES_DIR, "fig3_asymmetry_sweep.png"),
                      "Intermediate fraction vs asymmetry ratio (log scale). "
                      "Gold dashed line marks native Mce3R ratio (~20.4, Panagoda 2024).")

    # ── Slide 6: Figure 4 ────────────────────────────────────────────────
    make_figure_slide(prs, "Figure 4: Sensitivity Analysis",
                      os.path.join(FIGURES_DIR, "fig4_sensitivity.png"),
                      "Tornado plot of parameter impact on CV (left) and heatmap of "
                      "fold-change vs parameter showing number of modes (right).")

    # ── Slide 7: Figure 5 ────────────────────────────────────────────────
    make_figure_slide(prs, "Figure 5: Experimental Validation",
                      os.path.join(FIGURES_DIR, "fig5_validation.png"),
                      "Predicted vs published fold-change (A), persister fraction WT vs unregulated (B), "
                      "shuffle test (C), and symmetric TetR comparison (D).")

    # ── Slide 8: Figure 6 ────────────────────────────────────────────────
    make_figure_slide(prs, "Figure 6: Single-Cell Traces",
                      os.path.join(FIGURES_DIR, "fig6_single_cell_traces.png"),
                      "Time-series traces for 3 representative cells per condition. "
                      "~2100 min post-burn-in. Thin lines with transparency.")

    # ── Slide 9: Figure 7 ────────────────────────────────────────────────
    make_figure_slide(prs, "Figure 7: BIC Model Comparison",
                      os.path.join(FIGURES_DIR, "fig7_bic_comparison.png"),
                      "Grouped bar chart: BIC for 1-, 2-, and 3-component GMMs across conditions A\u2013D. "
                      "Star marks best (lowest) BIC per condition.")

    # ── Slide 10: Figure 8 ───────────────────────────────────────────────
    make_figure_slide(prs, "Figure 8: Methods Architecture",
                      os.path.join(FIGURES_DIR, "fig8_methods_architecture.png"),
                      "End-to-end computational pipeline: Phase 1 (bioinformatics) through "
                      "Phase 4 (figures). Color-coded by phase.")

    # ── Slide 11: Design Principles ──────────────────────────────────────
    slide = make_content_slide(prs, "Design Principles")
    items = [
        ("Consistent color scheme: ", "Teal = asymmetric, Orange = symmetric, Purple = single-site, Gray = unregulated"),
        ("Resolution: ", "All figures rendered at 300 dpi for publication quality"),
        ("Panel labels: ", "Bold (A), (B), (C)... in upper-left corner of each sub-panel"),
        ("Typography: ", "Calibri body, Trebuchet MS headers; 8\u201312 pt axis labels"),
        ("Shared style module: ", "figure_style.py enforces uniform aesthetics across all 8 figures"),
        ("Accessibility: ", "Color-blind-safe palette; patterns supplement color coding"),
        ("Format: ", "PNG output with transparent-free white backgrounds for journal submission"),
    ]
    add_bullets(slide, 0.6, 1.2, 12.0, 5.5, items, size=15, spacing=Pt(8))

    # ── Slide 12: Summary ────────────────────────────────────────────────
    make_summary_slide(prs, "Phase 4 Summary",
                       "8 publication-ready figures spanning binding site architecture\n"
                       "through statistical validation.\n\n"
                       "Figures 1\u20133: From genome-wide binding sites to asymmetry effects\n"
                       "Figures 4\u20135: Robustness checks and experimental grounding\n"
                       "Figures 6\u20137: Single-cell dynamics and statistical model selection\n"
                       "Figure 8: Complete methods architecture for reproducibility\n\n"
                       "All figures use a unified visual language defined in figure_style.py.")

    out = os.path.join(BASE, "Phase4_Presentation.pptx")
    prs.save(out)
    print(f"Saved Phase 4 presentation: {out}")


# ══════════════════════════════════════════════════════════════════════════
#  PHASE 7 PRESENTATION
# ══════════════════════════════════════════════════════════════════════════
def create_phase7():
    prs = new_prs()

    # ── Slide 1: Title ────────────────────────────────────────────────────
    make_title_slide(prs, "Phase 7: Extended Figures",
                     "ENIGMA Project \u2014 Visualizing Phases 5\u20136")

    # ── Slide 2: Figure Overview ──────────────────────────────────────────
    slide = make_content_slide(prs, "Extended Figure Overview")
    items = [
        ("Fig 9 \u2014 Repression Curves: ", "Fold-repression vs [Mce3R] for 4 operator architectures"),
        ("Fig 10 \u2014 Cooperativity Analysis: ", "MCMC posteriors + repression landscape heatmap"),
        ("Fig 11 \u2014 Environmental Distributions: ", "3\u00d74 grid of histograms across architectures and environments"),
        ("Fig 12 \u2014 Mutual Information: ", "MI (bits) vs concentration for 3 architectures"),
        ("Fig 13 \u2014 Persistence Phase Diagram: ", "Heatmap of persister fractions vs asymmetry ratio and environment"),
        ("Fig 14 \u2014 Two-Species Traces: ", "Mce3R + target protein dynamics with operator state coloring"),
    ]
    add_bullets(slide, 0.6, 1.2, 12.0, 5.5, items, size=16, spacing=Pt(12))

    # ── Slide 3: Figure 9 ────────────────────────────────────────────────
    make_figure_slide(prs, "Figure 9: Repression Curves",
                      os.path.join(EXT_FIGURES_DIR, "fig9_repression_curves.png"),
                      "Fold-repression vs [Mce3R] concentration for asymmetric, symmetric, "
                      "single-site, and unregulated architectures. Dashed line marks physiological concentration.")

    # ── Slide 4: Figure 10 ───────────────────────────────────────────────
    make_figure_slide(prs, "Figure 10: Cooperativity Analysis",
                      os.path.join(EXT_FIGURES_DIR, "fig10_cooperativity.png"),
                      "Panel A: 2D histogram of MCMC posteriors for log10(\u03c9) and \u0394G_spacer. "
                      "Panel B: Fold repression landscape at 332 nM Mce3R.")

    # ── Slide 5: Figure 11 ───────────────────────────────────────────────
    make_figure_slide(prs, "Figure 11: Environmental Distributions",
                      os.path.join(EXT_FIGURES_DIR, "fig11_environmental_distributions.png"),
                      "3\u00d74 grid: rows = asymmetric / symmetric / single-site; "
                      "columns = baseline / cholesterol / acidic pH / host-like. Shows how environment reshapes noise.")

    # ── Slide 6: Figure 12 ───────────────────────────────────────────────
    make_figure_slide(prs, "Figure 12: Mutual Information",
                      os.path.join(EXT_FIGURES_DIR, "fig12_mutual_information.png"),
                      "Mutual information (bits) vs [Mce3R] for 3 promoter architectures. "
                      "Vertical dashed line at physiological 332 nM.")

    # ── Slide 7: Figure 13 ───────────────────────────────────────────────
    make_figure_slide(prs, "Figure 13: Persistence Phase Diagram",
                      os.path.join(EXT_FIGURES_DIR, "fig13_persistence_phase.png"),
                      "Heatmap: persister fraction vs asymmetry ratio (X) and environment (Y). "
                      "Star marks the native Mce3R operator position (ratio ~20.4).")

    # ── Slide 8: Figure 14 ───────────────────────────────────────────────
    make_figure_slide(prs, "Figure 14: Two-Species Traces",
                      os.path.join(EXT_FIGURES_DIR, "fig14_two_species_traces.png"),
                      "Panel A: Time traces of Mce3R (blue) and target protein (orange) for 5 cells. "
                      "Panel B: Scatter of Mce3R vs target colored by operator state.")

    # ── Slide 9: The Extended Story ──────────────────────────────────────
    slide = make_content_slide(prs, "The Extended Story")
    items = [
        ("Thermodynamic foundations (Figs 9\u201310): ",
         "Repression curves establish the dose-response, while MCMC posteriors quantify cooperativity between binding sites"),
        ("Environmental modulation (Figs 11\u201312): ",
         "Protein distributions shift dramatically across environments; mutual information reveals which architectures transmit the most signal"),
        ("Persistence implications (Figs 13\u201314): ",
         "Phase diagram maps the landscape of persister fractions; two-species traces show how regulator-target dynamics co-evolve"),
        ("Narrative arc: ",
         "From equilibrium thermodynamics through stochastic environments to clinically relevant persistence phenotypes"),
    ]
    add_bullets(slide, 0.6, 1.2, 12.0, 5.5, items, size=15, spacing=Pt(10))

    # ── Slide 10: Summary ────────────────────────────────────────────────
    make_summary_slide(prs, "Phase 7 Summary",
                       "6 extended figures visualizing the results of Phases 5\u20136.\n\n"
                       "Figures 9\u201310: Thermodynamic repression and cooperativity\n"
                       "Figures 11\u201312: Environmental noise and information theory\n"
                       "Figures 13\u201314: Persistence phenotypes and two-species dynamics\n\n"
                       "Together with the 8 manuscript figures from Phase 4,\n"
                       "these 14 figures form a complete visual narrative of\n"
                       "Mce3R stochastic gene regulation in M. tuberculosis.")

    out = os.path.join(BASE, "Phase7_Presentation.pptx")
    prs.save(out)
    print(f"Saved Phase 7 presentation: {out}")


# ══════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    create_phase4()
    create_phase7()
    print("Done.")
